from openpyxl import load_workbook

from src.excel_exporter import run_report_pipeline, save_final_report


def test_report_contains_tables_charts_and_balanced_rows(tmp_path, simple_source_dataframe):
    source = tmp_path / "input.xlsx"
    output = tmp_path / "report.xlsx"
    simple_source_dataframe.to_excel(source, index=False)
    report = run_report_pipeline(source)
    save_final_report(report, output)
    workbook = load_workbook(output)
    try:
        assert len(workbook.sheetnames) == 11
        assert workbook.active.title == "Dashboard"
        assert len(workbook["Dashboard"]._charts) == 2
        assert "Development sample:" not in workbook["Dashboard"]["A30"].value
        assert workbook["Cleaned Transactions"].max_row == 3
        assert "CleanedTransactionsTable" in workbook["Cleaned Transactions"].tables
        assert len(workbook["Monthly Summary"].conditional_formatting) > 0
    finally:
        workbook.close()
