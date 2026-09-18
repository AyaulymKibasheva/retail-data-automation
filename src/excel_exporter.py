from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from time import perf_counter
from uuid import uuid4

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .analytics import AnalyticsResult, create_analytics
from .batch_loader import load_input_batch
from .cleaner import CleaningResult, clean_dataframe
from .transformer import TransformationResult, transform_cleaning_result
from .validator import ValidationReport, validate_dataframe


DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
DARK_GREY = "595959"
WHITE = "FFFFFF"
LIGHT_RED = "FCE8E6"
LIGHT_AMBER = "FFF2CC"
LIGHT_GREY = "E7E6E6"

TABLE_NAMES = {
    "Cleaned Transactions": "CleanedTransactionsTable",
    "Cancellations": "CancellationsTable",
    "Adjustments": "AdjustmentsTable",
    "Rejected Rows": "RejectedRowsTable",
    "Duplicates": "DuplicatesTable",
    "Monthly Summary": "MonthlySummaryTable",
    "Country Summary": "CountrySummaryTable",
    "Product Summary": "ProductSummaryTable",
    "Validation Issues": "ValidationIssuesTable",
}

THIN_GREY_BORDER = Border(
    bottom=Side(style="thin", color="D9E1F2")
)


class ExcelReportError(Exception):
    """Raised when the final Excel report cannot be created."""


@dataclass
class ReportData:
    raw_data: pd.DataFrame
    validation: ValidationReport
    cleaning: CleaningResult
    transformation: TransformationResult
    analytics: AnalyticsResult
    run_summary: pd.DataFrame
    stage_timings: pd.DataFrame


def measure_stage(records: list[dict], stage_name: str, function):
    started_at = perf_counter()
    result = function()
    records.append(
        {
            "stage": stage_name,
            "duration_seconds": round(perf_counter() - started_at, 4),
        }
    )
    return result


def run_report_pipeline(input_file: str | Path) -> ReportData:
    file_path = Path(input_file)
    if not file_path.exists():
        raise FileNotFoundError(f"Input file was not found: {file_path}")

    run_started_at = datetime.now()
    total_started_at = perf_counter()
    run_id = run_started_at.strftime("%Y%m%d-%H%M%S") + "-" + uuid4().hex[:8]
    stage_records: list[dict] = []

    input_batch = measure_stage(
        stage_records,
        "Discover, load and merge input files",
        lambda: load_input_batch(file_path),
    )
    raw_data = input_batch.dataframe
    validation = measure_stage(
        stage_records,
        "Validate data",
        lambda: validate_dataframe(raw_data),
    )

    if not validation.is_valid:
        errors = [issue.message for issue in validation.errors]
        raise ExcelReportError("Validation failed: " + " | ".join(errors))

    cleaning = measure_stage(
        stage_records,
        "Clean and classify data",
        lambda: clean_dataframe(raw_data),
    )
    transformation = measure_stage(
        stage_records,
        "Add business features",
        lambda: transform_cleaning_result(cleaning),
    )
    analytics = measure_stage(
        stage_records,
        "Calculate analytics",
        lambda: create_analytics(transformation),
    )

    run_finished_at = datetime.now()
    total_duration = perf_counter() - total_started_at
    classified_rows = (
        len(cleaning.clean_sales)
        + len(cleaning.cancellations)
        + len(cleaning.adjustments)
        + len(cleaning.rejected_rows)
        + len(cleaning.duplicates)
    )

    run_summary = pd.DataFrame(
        [
            {"property": "Run ID", "value": run_id},
            {"property": "Status", "value": "SUCCESS"},
            {"property": "Started at", "value": run_started_at},
            {"property": "Finished at", "value": run_finished_at},
            {
                "property": "Pipeline duration seconds",
                "value": round(total_duration, 4),
            },
            {"property": "Input path", "value": file_path.name},
            {"property": "Input files", "value": len(input_batch.files)},
            {
                "property": "Input size MB",
                "value": input_batch.total_size_mb,
            },
            {
                "property": "Combined input SHA256",
                "value": input_batch.combined_sha256,
            },
            {"property": "Rows loaded", "value": len(raw_data)},
            {"property": "Columns loaded", "value": raw_data.shape[1]},
            {"property": "Validation errors", "value": len(validation.errors)},
            {"property": "Validation warnings", "value": len(validation.warnings)},
            {"property": "Clean sales", "value": len(cleaning.clean_sales)},
            {"property": "Cancellations", "value": len(cleaning.cancellations)},
            {"property": "Adjustments", "value": len(cleaning.adjustments)},
            {"property": "Rejected rows", "value": len(cleaning.rejected_rows)},
            {"property": "Removed duplicates", "value": len(cleaning.duplicates)},
            {"property": "Classified rows", "value": classified_rows},
            {
                "property": "Row reconciliation",
                "value": classified_rows == len(raw_data),
            },
        ]
    )

    stage_records.append(
        {
            "stage": "Total pipeline",
            "duration_seconds": round(total_duration, 4),
        }
    )

    return ReportData(
        raw_data=raw_data,
        validation=validation,
        cleaning=cleaning,
        transformation=transformation,
        analytics=analytics,
        run_summary=run_summary,
        stage_timings=pd.DataFrame(stage_records),
    )


def style_header_row(worksheet, row_number: int = 1) -> None:
    for cell in worksheet[row_number]:
        cell.fill = PatternFill(fill_type="solid", fgColor=DARK_BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_GREY_BORDER
    worksheet.row_dimensions[row_number].height = 24


def set_column_widths(worksheet, maximum_width: int = 40) -> None:
    sample_end_row = min(worksheet.max_row, 201)
    for column_index in range(1, worksheet.max_column + 1):
        maximum_length = 0
        for row_index in range(1, sample_end_row + 1):
            value = worksheet.cell(row=row_index, column=column_index).value
            if value is not None:
                maximum_length = max(maximum_length, len(str(value)))
        column_letter = get_column_letter(column_index)
        worksheet.column_dimensions[column_letter].width = min(
            max(maximum_length + 2, 12), maximum_width
        )


def apply_column_formats(worksheet) -> None:
    currency_columns = {
        "revenue",
        "unit_price",
        "line_value",
        "cancellation_value",
        "adjustment_value",
        "average_order_value",
        "average_unit_price",
    }
    percentage_columns = {
        "revenue_share",
        "cancellation_rate",
        "anonymous_sales_share",
    }
    date_columns = {"invoice_date", "order_date"}
    integer_columns = {
        "quantity",
        "units_sold",
        "orders",
        "sales_rows",
        "registered_customers",
        "revenue_rank",
        "cancelled_orders",
        "cancelled_units",
        "cancellation_rows",
        "adjustment_rows",
    }
    headers = {cell.value: cell.column for cell in worksheet[1]}

    for header, column_index in headers.items():
        if header is None:
            continue
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if header in currency_columns:
                cell.number_format = "£#,##0.00"
            elif header in percentage_columns:
                cell.number_format = "0.00%"
            elif header in date_columns:
                cell.number_format = (
                    "yyyy-mm-dd hh:mm" if header == "invoice_date" else "yyyy-mm-dd"
                )
            elif header in integer_columns:
                cell.number_format = "#,##0"


def get_header_column(worksheet, header_name: str) -> int | None:
    for cell in worksheet[1]:
        if cell.value == header_name:
            return cell.column
    return None


def add_excel_table(worksheet, table_name: str) -> None:
    if worksheet.max_row < 2 or worksheet.max_column < 1:
        return

    table = Table(
        displayName=table_name,
        ref=(
            f"A1:{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        ),
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def add_data_bar(worksheet, header_name: str) -> None:
    column_index = get_header_column(worksheet, header_name)
    if column_index is None or worksheet.max_row < 2:
        return

    column_letter = get_column_letter(column_index)
    worksheet.conditional_formatting.add(
        f"{column_letter}2:{column_letter}{worksheet.max_row}",
        DataBarRule(
            start_type="min",
            end_type="max",
            color=MEDIUM_BLUE,
            showValue=True,
        ),
    )


def add_row_highlight(
    worksheet,
    header_name: str,
    formula_condition: str,
    fill_color: str,
    font_color: str = DARK_GREY,
    stop_if_true: bool = False,
) -> None:
    column_index = get_header_column(worksheet, header_name)
    if column_index is None or worksheet.max_row < 2:
        return

    column_letter = get_column_letter(column_index)
    data_range = (
        f"A2:{get_column_letter(worksheet.max_column)}"
        f"{worksheet.max_row}"
    )
    formula = formula_condition.format(column=f"${column_letter}2")
    worksheet.conditional_formatting.add(
        data_range,
        FormulaRule(
            formula=[formula],
            fill=PatternFill(fill_type="solid", fgColor=fill_color),
            font=Font(color=font_color),
            stopIfTrue=stop_if_true,
        ),
    )


def apply_conditional_formatting(worksheet) -> None:
    sheet_name = worksheet.title

    if sheet_name in {
        "Monthly Summary",
        "Country Summary",
        "Product Summary",
    }:
        add_data_bar(worksheet, "revenue")

    if sheet_name == "Cleaned Transactions":
        add_row_highlight(
            worksheet,
            "data_quality_status",
            'ISNUMBER(SEARCH("missing_description",{column}))',
            LIGHT_RED,
            "9C0006",
            stop_if_true=True,
        )
        add_row_highlight(
            worksheet,
            "data_quality_status",
            'ISNUMBER(SEARCH("missing_customer",{column}))',
            LIGHT_AMBER,
            "7F6000",
        )
    elif sheet_name == "Cancellations":
        add_row_highlight(
            worksheet,
            "record_type",
            '{column}="cancellation"',
            LIGHT_RED,
            "9C0006",
        )
    elif sheet_name == "Adjustments":
        add_row_highlight(
            worksheet,
            "adjustment_reason",
            "NOT(ISBLANK({column}))",
            LIGHT_AMBER,
            "7F6000",
        )
    elif sheet_name == "Rejected Rows":
        add_row_highlight(
            worksheet,
            "rejection_reason",
            "NOT(ISBLANK({column}))",
            LIGHT_RED,
            "9C0006",
        )
    elif sheet_name == "Duplicates":
        add_row_highlight(
            worksheet,
            "record_type",
            '{column}="duplicate"',
            LIGHT_GREY,
        )
    elif sheet_name == "Validation Issues":
        add_row_highlight(
            worksheet,
            "level",
            '{column}="error"',
            LIGHT_RED,
            "9C0006",
            stop_if_true=True,
        )
        add_row_highlight(
            worksheet,
            "level",
            '{column}="warning"',
            LIGHT_AMBER,
            "7F6000",
            stop_if_true=True,
        )
        add_row_highlight(
            worksheet,
            "level",
            '{column}="info"',
            LIGHT_BLUE,
            DARK_BLUE,
        )


def style_tabular_sheet(worksheet, table_name: str) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    style_header_row(worksheet)
    if worksheet.max_row >= 2:
        add_excel_table(worksheet, table_name)
    elif worksheet.max_row == 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    set_column_widths(worksheet)
    apply_column_formats(worksheet)
    apply_conditional_formatting(worksheet)


def write_kpi_card(
    worksheet,
    start_column: int,
    label_row: int,
    label: str,
    value,
    number_format: str,
) -> None:
    end_column = start_column + 2
    worksheet.merge_cells(
        start_row=label_row,
        start_column=start_column,
        end_row=label_row,
        end_column=end_column,
    )
    worksheet.merge_cells(
        start_row=label_row + 1,
        start_column=start_column,
        end_row=label_row + 2,
        end_column=end_column,
    )

    label_cell = worksheet.cell(row=label_row, column=start_column)
    value_cell = worksheet.cell(row=label_row + 1, column=start_column)
    label_cell.value = label
    label_cell.fill = PatternFill(fill_type="solid", fgColor=MEDIUM_BLUE)
    label_cell.font = Font(color=WHITE, bold=True, size=11)
    label_cell.alignment = Alignment(horizontal="center", vertical="center")

    value_cell.value = value
    value_cell.fill = PatternFill(fill_type="solid", fgColor=LIGHT_BLUE)
    value_cell.font = Font(color=DARK_BLUE, bold=True, size=18)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.number_format = number_format

    card_border = Border(
        left=Side(style="thin", color="B4C6E7"),
        right=Side(style="thin", color="B4C6E7"),
        top=Side(style="thin", color="B4C6E7"),
        bottom=Side(style="thin", color="B4C6E7"),
    )
    for row in range(label_row, label_row + 3):
        for column in range(start_column, end_column + 1):
            worksheet.cell(row=row, column=column).border = card_border


def create_dashboard(workbook, analytics: AnalyticsResult) -> None:
    worksheet = workbook["Dashboard"]
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:L2")
    title_cell = worksheet["A1"]
    title_cell.value = "Retail Data Automation Report"
    title_cell.fill = PatternFill(fill_type="solid", fgColor=DARK_BLUE)
    title_cell.font = Font(color=WHITE, bold=True, size=22)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.merge_cells("A3:L3")
    subtitle_cell = worksheet["A3"]
    subtitle_cell.value = "Sales, cancellations and data-quality summary"
    subtitle_cell.font = Font(color=DARK_GREY, italic=True, size=11)
    subtitle_cell.alignment = Alignment(horizontal="center")

    kpi_values = analytics.kpis.set_index("metric")["value"].to_dict()
    cards = [
        (1, 5, "Total revenue", kpi_values["Total revenue"], "£#,##0.00"),
        (4, 5, "Sales orders", kpi_values["Sales orders"], "#,##0"),
        (7, 5, "Average order value", kpi_values["Average order value"], "£#,##0.00"),
        (10, 5, "Units sold", kpi_values["Units sold"], "#,##0"),
        (1, 9, "Registered customers", kpi_values["Unique registered customers"], "#,##0"),
        (4, 9, "Cancellation rate", kpi_values["Cancellation rate"], "0.00%"),
        (7, 9, "Cancellation value", kpi_values["Cancellation value"], "£#,##0.00"),
        (10, 9, "Countries", kpi_values["Countries"], "#,##0"),
    ]
    for card in cards:
        write_kpi_card(worksheet, *card)

    monthly_sheet = workbook["Monthly Summary"]
    monthly_chart = LineChart()
    monthly_chart.title = "Monthly Revenue"
    monthly_chart.style = 10
    monthly_chart.height = 7
    monthly_chart.width = 13
    monthly_chart.legend = None
    monthly_chart.varyColors = False
    monthly_chart.visible_cells_only = False
    monthly_chart.x_axis.delete = False
    monthly_chart.x_axis.tickLblPos = "low"
    monthly_chart.x_axis.tickLblSkip = 1
    monthly_chart.y_axis.delete = False
    monthly_chart.y_axis.tickLblPos = "nextTo"

    monthly_chart.add_data(
        Reference(
            monthly_sheet,
            min_col=2,
            min_row=1,
            max_row=monthly_sheet.max_row,
        ),
        titles_from_data=True,
    )
    monthly_chart.set_categories(
        Reference(
            monthly_sheet,
            min_col=1,
            min_row=2,
            max_row=monthly_sheet.max_row,
        )
    )
    if monthly_chart.series:
        series = monthly_chart.series[0]
        series.graphicalProperties.line.solidFill = MEDIUM_BLUE
        series.graphicalProperties.line.width = 28575
        series.smooth = False
        series.marker.symbol = "circle"
        series.marker.size = 6
        series.marker.graphicalProperties.solidFill = MEDIUM_BLUE
        series.marker.graphicalProperties.line.solidFill = DARK_BLUE
    worksheet.add_chart(monthly_chart, "A14")

    international_countries = (
        analytics.country_summary.loc[
            analytics.country_summary["country"].ne("United Kingdom")
        ]
        .head(8)
        .copy()
        .sort_values("revenue", ascending=True)
    )
    worksheet["N1"] = "Country"
    worksheet["O1"] = "Revenue"
    for cell in (worksheet["N1"], worksheet["O1"]):
        cell.fill = PatternFill(fill_type="solid", fgColor=DARK_BLUE)
        cell.font = Font(color=WHITE, bold=True)

    for excel_row, record in enumerate(
        international_countries.itertuples(index=False), start=2
    ):
        worksheet.cell(row=excel_row, column=14, value=record.country)
        revenue_cell = worksheet.cell(
            row=excel_row,
            column=15,
            value=float(record.revenue),
        )
        revenue_cell.number_format = "£#,##0.00"

    country_chart = BarChart()
    country_chart.type = "bar"
    country_chart.style = 10
    country_chart.title = "International Revenue by Country"
    country_chart.height = 7.8
    country_chart.width = 15.5
    country_chart.legend = None
    country_chart.varyColors = False
    country_chart.visible_cells_only = False
    country_chart.gapWidth = 60
    country_chart.x_axis.delete = False
    country_chart.x_axis.tickLblPos = "low"
    country_chart.y_axis.delete = False
    country_chart.y_axis.tickLblPos = "nextTo"
    country_chart.y_axis.majorUnit = 1000
    country_chart.y_axis.numFmt = '£0,"k"'
    country_chart.dLbls = DataLabelList()
    country_chart.dLbls.showVal = True
    country_chart.dLbls.showSerName = False
    country_chart.dLbls.showCatName = False
    country_chart.dLbls.showLegendKey = False
    country_chart.dLbls.showPercent = False
    country_chart.dLbls.showBubbleSize = False
    country_chart.dLbls.numFmt = "£#,##0"
    country_chart.dLbls.dLblPos = "outEnd"
    country_end_row = len(international_countries) + 1
    country_chart.add_data(
        Reference(
            worksheet,
            min_col=15,
            min_row=1,
            max_row=country_end_row,
        ),
        titles_from_data=True,
    )
    country_chart.set_categories(
        Reference(
            worksheet,
            min_col=14,
            min_row=2,
            max_row=country_end_row,
        )
    )
    if country_chart.series:
        series = country_chart.series[0]
        series.graphicalProperties.solidFill = MEDIUM_BLUE
        series.graphicalProperties.line.solidFill = DARK_BLUE
    worksheet.add_chart(country_chart, "G14")

    worksheet.merge_cells("A30:L30")
    note_cell = worksheet["A30"]
    note_cell.value = "Development sample: results are not full-store business totals."
    note_cell.fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    note_cell.font = Font(color="7F6000", italic=True)
    note_cell.alignment = Alignment(horizontal="center")

    worksheet.merge_cells("A32:L32")
    source_cell = worksheet["A32"]
    source_cell.value = (
        "Source: Chen, D. (2015), Online Retail, "
        "UCI Machine Learning Repository, DOI 10.24432/C5BW33"
    )
    source_cell.font = Font(color=DARK_GREY, size=9)
    source_cell.alignment = Alignment(horizontal="left")

    for column_index in range(1, 13):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 14
    worksheet.column_dimensions["N"].hidden = True
    worksheet.column_dimensions["O"].hidden = True
    worksheet.sheet_view.selection[0].activeCell = "A1"
    worksheet.sheet_view.selection[0].sqref = "A1"


def style_processing_log(
    worksheet,
    run_summary_rows: int,
    stage_timing_rows: int,
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A3"
    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "Processing run"
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor=DARK_BLUE)
    worksheet["A1"].font = Font(color=WHITE, bold=True, size=14)
    worksheet["A1"].alignment = Alignment(horizontal="center")
    style_header_row(worksheet, row_number=2)

    stage_title_row = run_summary_rows + 4
    worksheet.merge_cells(
        start_row=stage_title_row,
        start_column=1,
        end_row=stage_title_row,
        end_column=2,
    )
    stage_title_cell = worksheet.cell(row=stage_title_row, column=1)
    stage_title_cell.value = "Stage timings"
    stage_title_cell.fill = PatternFill(fill_type="solid", fgColor=DARK_BLUE)
    stage_title_cell.font = Font(color=WHITE, bold=True, size=12)
    style_header_row(worksheet, row_number=stage_title_row + 1)
    set_column_widths(worksheet, maximum_width=70)

    run_table = Table(
        displayName="ProcessingRunTable",
        ref=f"A2:B{run_summary_rows + 2}",
    )
    run_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(run_table)

    if stage_timing_rows:
        timing_header_row = stage_title_row + 1
        timing_table = Table(
            displayName="StageTimingsTable",
            ref=(
                f"A{timing_header_row}:"
                f"B{timing_header_row + stage_timing_rows}"
            ),
        )
        timing_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(timing_table)


def save_final_report(report: ReportData, output_file: str | Path) -> None:
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            workbook = writer.book
            dashboard = workbook.create_sheet("Dashboard")
            writer.sheets["Dashboard"] = dashboard

            exports = {
                "Cleaned Transactions": report.transformation.sales,
                "Cancellations": report.transformation.cancellations,
                "Adjustments": report.transformation.adjustments,
                "Rejected Rows": report.transformation.rejected_rows,
                "Duplicates": report.transformation.duplicates,
                "Monthly Summary": report.analytics.monthly_summary,
                "Country Summary": report.analytics.country_summary,
                "Product Summary": report.analytics.product_summary,
            }
            for sheet_name, dataframe in exports.items():
                dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

            report.run_summary.to_excel(
                writer,
                sheet_name="Processing Log",
                startrow=1,
                index=False,
            )
            stage_title_row = len(report.run_summary) + 4
            report.stage_timings.to_excel(
                writer,
                sheet_name="Processing Log",
                startrow=stage_title_row,
                index=False,
            )

            report.validation.to_dataframe().to_excel(
                writer,
                sheet_name="Validation Issues",
                index=False,
            )

            for sheet_name in exports:
                style_tabular_sheet(
                    workbook[sheet_name],
                    TABLE_NAMES[sheet_name],
                )
            style_tabular_sheet(
                workbook["Validation Issues"],
                TABLE_NAMES["Validation Issues"],
            )
            style_processing_log(
                workbook["Processing Log"],
                len(report.run_summary),
                len(report.stage_timings),
            )
            create_dashboard(workbook, report.analytics)
            workbook.active = workbook.sheetnames.index("Dashboard")
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True

    except PermissionError as error:
        raise ExcelReportError(
            f"Cannot save '{output_path.name}'. "
            "Close the file in Excel and run again."
        ) from error


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    input_file = project_root / "data" / "sample" / "sample_transactions.xlsx"
    output_file = project_root / "output" / "retail_automation_report.xlsx"

    print("Running report pipeline...")
    report = run_report_pipeline(input_file)
    print("Creating formatted Excel report...")
    save_final_report(report, output_file)
    print("\nREPORT SUMMARY")
    print(report.analytics.kpis.to_string(index=False))
    print(f"\nReport saved to: {output_file}")
    print("Excel report created successfully.")


if __name__ == "__main__":
    main()
